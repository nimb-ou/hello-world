#!/usr/bin/env python3
"""
End-to-end pipeline for the LORE financial model.

1. build_model.py         -> raw workbook with live formulas (no cached values)
2. LibreOffice recalc     -> bear / bull scenario values (snapshot)
3. write snapshot table into Scenarios sheet, recalc base -> finance/LORE_Financial_Model.xlsx
   (formulas + cached values, so it opens with numbers in any viewer)
4. verify: balance sheet balances, cash ties, independent Python re-computation of
   units / hardware revenue / subscriber roll-forward / revenue totals
5. export finance/model_summary.json and charts (finance/charts/*.png)

Run: python3 finance/scripts/run_all.py
Requires LibreOffice Calc (apt-get install libreoffice-calc) and openpyxl, matplotlib.
"""
import json
import os
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.dirname(HERE)
FINAL = os.path.join(FIN, "LORE_Financial_Model.xlsx")
TMP = tempfile.mkdtemp(prefix="lore_fin_")
PY = sys.executable


def sh(*args, env=None):
    subprocess.run(list(args), check=True, env=env)


def recalc(src, dst, scen=None):
    a = [PY, os.path.join(HERE, "lo_recalc.py"), src, dst]
    if scen:
        a.append(str(scen))
    sh(*a)


# 1. build raw
raw = os.path.join(TMP, "raw.xlsx")
sh(PY, os.path.join(HERE, "build_model.py"), env={**os.environ, "LORE_OUT": raw})
RM = json.load(open(os.path.join(HERE, "rowmap.json")))
SUMC, COLS, FY = RM["sumcols"], RM["cols"], RM["fy"]


def metrics(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    S = RM["sheets"]

    def fy(sheet, key):
        ws = wb[sheet]
        return [ws[f"{c}{S[sheet][key]}"].value for c in SUMC]

    def per(sheet, key):
        ws = wb[sheet]
        return [ws[f"{c}{S[sheet][key]}"].value for c in COLS]

    fr = wb["Funding & Runway"]
    met = {k: fr[f"D{r}"].value for k, r in RM["fr_metrics"].items()}
    return wb, fy, per, met


# 2. scenarios
snap = {}
for scen, name in [(1, "Bear"), (3, "Bull"), (2, "Base")]:
    out = os.path.join(TMP, f"s{scen}.xlsx")
    recalc(raw, out, scen)
    wb, fy, per, met = metrics(out)
    snap[name] = dict(rev=fy("P&L", "rev"), gm=fy("P&L", "gm"), ebitda=fy("P&L", "ebitda"),
                      cash=fy("Cash Flow", "end"), units=fy("P&L", "units"), arr=fy("P&L", "arr"),
                      min_cash=met["Minimum period-end cash incl. FY29-31"],
                      need=met["Peak cumulative op+inv cash outflow (= minimum equity need before buffer)"],
                      be=met["First fiscal year with positive EBITDA"])

# 3. snapshot into Scenarios sheet of the raw (formula) workbook, then recalc base
wb = openpyxl.load_workbook(raw)
ws = wb["Scenarios"]
r = RM["sc_snap_row"]
ws.cell(r, 1, "SCENARIO SNAPSHOT (static values written by scripts/run_all.py after LibreOffice recalc of each scenario; "
              "re-run the script after changing inputs)").font = Font(bold=True)
hdr = ["Metric", "Scenario"] + [f"FY{y}" for y in FY]
for j, h in enumerate(hdr):
    c = ws.cell(r + 1, 1 + j, h)
    c.font, c.fill = Font(bold=True, color="F3EEE6"), PatternFill("solid", fgColor="14161A")
rr = r + 2
for lab, key, fmt in [("Units shipped", "units", "#,##0"), ("Revenue", "rev", "#,##0"), ("Gross margin %", "gm", "0.0%"),
                      ("EBITDA", "ebitda", "#,##0"), ("Ending cash", "cash", "#,##0"), ("Subscription ARR (end)", "arr", "#,##0")]:
    for name in ["Bear", "Base", "Bull"]:
        ws.cell(rr, 1, lab)
        ws.cell(rr, 2, name)
        for k, v in enumerate(snap[name][key]):
            c = ws.cell(rr, 3 + k, v)
            c.number_format = fmt
        rr += 1
for name in ["Bear", "Base", "Bull"]:
    ws.cell(rr, 1, "Min cash (with planned rounds) / peak cumulative burn / first EBITDA+ FY")
    ws.cell(rr, 2, name)
    ws.cell(rr, 3, snap[name]["min_cash"]).number_format = "#,##0"
    ws.cell(rr, 4, snap[name]["need"]).number_format = "#,##0"
    ws.cell(rr, 5, snap[name]["be"])
    rr += 1
raw2 = os.path.join(TMP, "raw2.xlsx")
wb.save(raw2)
recalc(raw2, FINAL, 2)

# 4. verification
wb, fy, per, met = metrics(FINAL)
errs = []
for sname in wb.sheetnames:
    for row in wb[sname].iter_rows():
        for c in row:
            if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                errs.append(f"{sname}!{c.coordinate}={c.value}")
checks = []
chk = per("Balance Sheet", "check")
checks.append(("Balance sheet balances every period", all(abs(v or 0) < 1 for v in chk)))
cash_bs, cash_cf = per("Balance Sheet", "cash"), per("Cash Flow", "end")
checks.append(("Cash flow ending cash == balance sheet cash", all(abs(a - b) < 1 for a, b in zip(cash_bs, cash_cf))))
checks.append(("No formula errors in workbook", not errs))

# independent python recompute (monthly periods) from the input cells
A = wb["Assumptions"]
names = {n: wb.defined_names[n].attr_text for n in wb.defined_names.keys()}


def val(n):
    sheet, cell = names[n].split("!")
    return wb[sheet.strip("'")][cell.replace("$", "")].value


TS = {k: [A[f"{c}{r}"].value for c in COLS] for k, r in RM["tsin"].items()}
P = RM["periods"]
NM = sum(p["monthly"] for p in P)
vol, price, att_m, ch_m = val("sc_vol"), val("sc_price"), val("sc_attach"), val("sc_churn")
pre_tot = (val("ks_units") + val("pre_units")) * vol
mixp = {"peb": val("pre_mix_peb"), "atl": val("pre_mix_atl")}
mixp["sq"] = 1 - mixp["peb"] - mixp["atl"]
units = {s: [] for s in ("peb", "sq", "atl")}
hw_rev = []
for i in range(len(P)):
    reg = TS["reg_units"][i] * vol
    rs = TS["retail_share"][i]
    rev = 0
    for s in ("peb", "sq", "atl"):
        m = TS[f"mix_{s}"][i]
        msrp = val(f"msrp_{s}") * price
        ship = TS["pre_ship_atl"][i] if s == "atl" else TS["pre_ship_pct"][i]
        pre_u = pre_tot * ship * mixp[s]
        pre_px = (val("ks_units") * val(f"ks_{s}") + val("pre_units") * val(f"pre_{s}")) / (val("ks_units") + val("pre_units")) * price
        u = reg * m + pre_u
        units[s].append(u)
        rev += reg * m * (1 - rs) * msrp * (1 - val("dtc_disc")) + reg * m * rs * msrp * (1 - val("retail_margin")) + pre_u * pre_px
    hw_rev.append(rev * (1 - val("returns_pct")))
model_hw = per("Revenue Build", "hw_net")
checks.append(("Independent recompute: net hardware revenue (all periods)",
               all(abs(a - b) < 1 for a, b in zip(hw_rev, model_hw))))
# subscriber roll-forward (monthly)
churn = min(0.5, val("churn") * ch_m)
conv = [sum(units[s][i] * min(1, val(f"attach_{s}") * att_m) for s in units) for i in range(len(P))]
end, subs_ok = 0.0, True
model_end = per("Subscription Cohorts", "end")
for i in range(NM):
    new = conv[i - 3] if i >= 3 else 0
    end = end * (1 - churn) + new
    subs_ok &= abs(end - model_end[i]) < 0.01
checks.append(("Independent recompute: monthly paid-subscriber roll-forward", subs_ok))
tot = per("P&L", "rev")
parts = [per("P&L", k) for k in ("hw_net", "subs", "trip", "acc", "aff", "fleet")]
checks.append(("Revenue streams sum to total revenue", all(abs(sum(p[i] for p in parts) - tot[i]) < 1 for i in range(len(P)))))
ct = wb["Cap Table"]
tot_pct = [ct.cell(RM["cap"]["tot"], 3 + 2 * j).value for j in range(5)]
checks.append(("Cap table sums to 100% at every stage", all(abs(v - 1) < 1e-9 for v in tot_pct)))
fy_rev_sum = sum(fy("P&L", "rev"))
checks.append(("Annual summary == sum of all periods (revenue)", abs(fy_rev_sum - sum(tot)) < 1))

print("\nVERIFICATION")
for lab, ok in checks:
    print(("PASS " if ok else "FAIL ") + lab)
if errs:
    print("errors:", errs[:10])

# 5. export summary JSON
S = RM["sheets"]
summary = dict(
    fy=FY,
    scenario_snapshot=snap,
    checks={lab: bool(ok) for lab, ok in checks},
    annual={k: fy(sh_, key) for k, (sh_, key) in {
        "units": ("P&L", "units"), "revenue": ("P&L", "rev"), "hw_net": ("P&L", "hw_net"), "subs": ("P&L", "subs"),
        "trip": ("P&L", "trip"), "acc": ("P&L", "acc"), "aff": ("P&L", "aff"), "fleet": ("P&L", "fleet"),
        "cogs": ("P&L", "cogs"), "gp": ("P&L", "gp"), "gm": ("P&L", "gm"), "hw_gm": ("P&L", "hw_gm"),
        "rd": ("P&L", "rd"), "sm": ("P&L", "sm"), "ops": ("P&L", "ops"), "ga": ("P&L", "ga"), "opex": ("P&L", "opex"),
        "ebitda": ("P&L", "ebitda"), "ebitda_m": ("P&L", "ebitda_m"), "ni": ("P&L", "ni"),
        "cash_end": ("Cash Flow", "end"), "burn": ("Cash Flow", "burn"), "equity": ("Cash Flow", "equity"),
        "subs_end": ("Subscription Cohorts", "end"), "arr": ("Subscription Cohorts", "arr"),
        "active_end": ("Subscription Cohorts", "act_end"), "hc": ("Headcount Plan", "hc_total"),
        "hc_rd": ("Headcount Plan", "hcd_R&D"), "hc_sm": ("Headcount Plan", "hcd_S&M"),
        "hc_ops": ("Headcount Plan", "hcd_Ops & CS"), "hc_ga": ("Headcount Plan", "hcd_G&A"),
        "cloud_cogs": ("Cloud & AI COGS", "c_total"), "nre": ("OpEx", "nre_total"), "capex": ("Working Capital", "capex"),
        "inv_end": ("Working Capital", "inv"), "sub_rev_total": ("Revenue Build", "recurring"),
    }.items()},
    periods=[p["label"] for p in P],
    monthly={k: per(sh_, key) for k, (sh_, key) in {
        "cash": ("Cash Flow", "end"), "burn": ("Cash Flow", "burn"), "revenue": ("P&L", "rev"), "ebitda": ("P&L", "ebitda"),
        "units": ("P&L", "units"), "hc": ("Headcount Plan", "hc_total"), "runway": ("Runway", "runway"),
        "rd": ("P&L", "rd"), "sm": ("P&L", "sm"), "ops": ("P&L", "ops"), "ga": ("P&L", "ga"),
        "rd_pc": ("OpEx", "rd_pc"), "nre_id": ("OpEx", "nre_id"), "nre_odm": ("OpEx", "nre_odm"), "nre_proto": ("OpEx", "nre_proto"),
        "nre_cert": ("OpEx", "nre_cert"), "nre_content": ("OpEx", "nre_content"), "nre_sw": ("OpEx", "nre_sw"),
        "sm_pc": ("OpEx", "sm_pc"), "ks_mkt": ("OpEx", "ks_mkt"), "ks_fees": ("OpEx", "ks_fees"), "brand": ("OpEx", "brand"),
        "paid": ("OpEx", "paid"), "ga_pc": ("OpEx", "ga_pc"), "ga_other": ("OpEx", "ga_other"), "ops_pc": ("OpEx", "ops"),
        "capex": ("Working Capital", "capex"), "purch": ("Working Capital", "purch"), "pre_cash": ("Working Capital", "pre_cash"),
        "equity": ("Cash Flow", "equity"), "cfo": ("Cash Flow", "cfo"), "gp": ("P&L", "gp"), "deferred": ("Working Capital", "deferred"),
    }.items()},
    funding=met,
)
ue = wb["Unit Economics"]
summary["unit_economics"] = {ue.cell(r, 1).value: [ue.cell(r, 2 + j).value for j in range(6)] for r in RM["ue_labels"].values()}
summary["cap_table"] = [[ct.cell(RM["cap"]["hr0"] + k, 1).value] + [ct.cell(RM["cap"]["hr0"] + k, 3 + 2 * j).value for j in range(5)]
                        for k in range(RM["cap"]["tot"] - RM["cap"]["hr0"])]
json.dump(summary, open(os.path.join(FIN, "model_summary.json"), "w"), indent=1, default=str)
sh(PY, os.path.join(HERE, "make_charts.py"))
print("done ->", FINAL)
if not all(ok for _, ok in checks):
    sys.exit(1)
