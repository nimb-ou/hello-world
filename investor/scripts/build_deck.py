#!/usr/bin/env python3
"""Build investor/LORE_Pitch_Deck.pptx (16:9, 16 slides) from the teams' work.

Run from anywhere:  python3 investor/scripts/build_deck.py
Needs: python-pptx, Pillow. Finance numbers are read live from finance/model_summary.json
when it exists; otherwise the financial slides show clearly marked placeholders.
"""
import json
import os
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
INV = ROOT / "investor"
ASSETS = INV / "assets"
ASSETS.mkdir(exist_ok=True)
OUT = INV / "LORE_Pitch_Deck.pptx"

# ---------------------------------------------------------------- brand
INK = RGBColor(0x14, 0x16, 0x1A)
BONE = RGBColor(0xF3, 0xEE, 0xE6)
TERRA = RGBColor(0xC4, 0x62, 0x3A)
BRASS = RGBColor(0xB8, 0x93, 0x5A)
OLIVE = RGBColor(0x5E, 0x6B, 0x4E)
FOG = RGBColor(0x9A, 0xA0, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFB, 0xF8, 0xF3)
LINE = RGBColor(0xDD, 0xD5, 0xC8)
MUTED = RGBColor(0x5A, 0x5D, 0x63)
# Fraunces / Inter are the brand fonts (OFL, brand/typography/fonts). If they are not
# installed on the viewing machine PowerPoint substitutes (Georgia / Arial-like).
DISPLAY = "Fraunces"
BODY = "Inter"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)  # side margin

# ---------------------------------------------------------------- finance
FIN_JSON = ROOT / "finance" / "model_summary.json"
FIN = json.loads(FIN_JSON.read_text()) if FIN_JSON.exists() else None


def money(v, dp=1):
    if v is None:
        return "[TBD]"
    a = abs(v)
    s = "-" if v < 0 else ""
    if a >= 1e9:
        return f"{s}${a/1e9:.{dp}f}B"
    if a >= 1e6:
        return f"{s}${a/1e6:.{dp}f}M"
    if a >= 1e3:
        return f"{s}${a/1e3:.0f}k"
    return f"{s}${a:.0f}"


# ---------------------------------------------------------------- assets
def prep_assets():
    """Crop renders/screenshots so stale captions (e.g. the $549 Atlas label on
    lineup.png, superseded by D8 -> $599) are not shown."""
    r = ROOT / "design/industrial/renders"
    s = ROOT / "software/apps/prototype/screenshots"
    jobs = {
        "lineup.png": (r / "lineup.png", (150, 420, 2300, 1080)),
        "hero-pebble.png": (r / "hero-pebble.png", (250, 120, 2150, 1440)),
        "hero-square.png": (r / "hero-square.png", (250, 120, 2150, 1440)),
        "hero-atlas.png": (r / "hero-atlas.png", (250, 120, 2150, 1440)),
        "exploded-square.png": (r / "exploded-square.png", (0, 0, 2400, 1460)),
        "shot-pantheon.png": (s / "01-sim-pantheon.png", (0, 0, 780, 1600)),
        "shot-navona.png": (s / "02-sim-piazza-navona.png", (0, 0, 780, 1600)),
        "shot-around.png": (s / "04-sim-whats-around-me-largo-argentina.png", (0, 0, 780, 1600)),
        "shot-dark.png": (s / "08-sim-dark-storyteller.png", (0, 0, 780, 1600)),
    }
    out = {}
    for name, (src, box) in jobs.items():
        if not src.exists():
            continue
        im = Image.open(src)
        x0, y0, x1, y1 = box
        im = im.crop((x0, y0, min(x1, im.width), min(y1, im.height)))
        p = ASSETS / name
        im.save(p)
        out[name] = p
    return out


A = prep_assets()
LOGO_INK = ROOT / "brand/logo/png/lore-lockup-horizontal-ink-1600w.png"
LOGO_BONE = ROOT / "brand/logo/png/lore-lockup-horizontal-bone-1600w.png"
SYMBOL_T = ROOT / "brand/logo/png/lore-symbol-terracotta-512.png"
WORD_BONE = ROOT / "brand/logo/png/lore-wordmark-bone-1600w.png"
RCH = ROOT / "research/charts"
FCH = ROOT / "finance/charts"

# ---------------------------------------------------------------- helpers
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
SLIDE_NO = [0]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1, space_after=4):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        segs = para if isinstance(para, list) else [(para, {})]
        for seg in segs:
            if isinstance(seg, str):
                seg = (seg, {})
            t, o = seg
            r = p.add_run()
            r.text = t
            r.font.size = Pt(o.get("size", size))
            r.font.name = o.get("font", font)
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def image(slide, path, x, y, w, h, align="center"):
    """Place image contained within the box, keeping aspect ratio."""
    path = Path(path)
    if not path.exists():
        box(slide, x, y, w, h, fill=CARD, line=LINE)
        text(slide, x, y, w, h, f"[missing: {path.name}]", size=11, color=MUTED,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return None
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px = x + (w - pw) // 2 if align == "center" else x
    py = y + (h - ph) // 2
    return slide.shapes.add_picture(str(path), px, py, pw, ph)


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    bg(s, INK if dark else BONE)
    SLIDE_NO[0] += 1
    return s


def header(s, kicker, title, dark=False, sub=None):
    fg = BONE if dark else INK
    text(s, MX, Inches(0.42), Inches(9), Inches(0.3), kicker.upper(), size=11,
         color=TERRA, bold=True)
    tsize = 30 if (len(title) <= 54 or "\n" in title) else 26
    text(s, MX, Inches(0.72), Inches(12.1), Inches(0.9), title, size=tsize, font=DISPLAY,
         color=fg, spacing=1.0)
    if sub:
        text(s, MX, Inches(1.5), Inches(12.1), Inches(0.4), sub, size=13,
             color=FOG if dark else MUTED)
    footer(s, dark)


def footer(s, dark=False):
    logo = LOGO_BONE if dark else LOGO_INK
    if logo.exists():
        image(s, logo, MX, Inches(7.02), Inches(0.9), Inches(0.24), align="left")
    text(s, Inches(2.0), Inches(7.06), Inches(8), Inches(0.25),
         "LORE Labs · Pre-seed 2026 · Confidential", size=9, color=FOG)
    text(s, Inches(11.9), Inches(7.06), Inches(0.83), Inches(0.25), f"{SLIDE_NO[0]:02d}",
         size=9, color=FOG, align=PP_ALIGN.RIGHT)


def source(s, txt, dark=False):
    text(s, MX, Inches(6.68), Inches(12.1), Inches(0.3), txt, size=8.5,
         color=FOG if dark else MUTED)


def card(s, x, y, w, h, title=None, body=None, fill=CARD, tcolor=INK, bcolor=MUTED,
         tsize=15, bsize=11.5, accent=TERRA, line=LINE):
    box(s, x, y, w, h, fill=fill, line=line)
    if accent is not None:
        box(s, x, y, Inches(0.06), h, fill=accent)
    pad = Inches(0.22)
    yy = y + Inches(0.16)
    if title:
        text(s, x + pad, yy, w - 2 * pad, Inches(0.5), title, size=tsize, bold=True,
             color=tcolor)
        yy += Inches(0.42 if tsize <= 16 else 0.6)
    if body:
        text(s, x + pad, yy, w - 2 * pad, h - (yy - y) - Inches(0.1), body, size=bsize,
             color=bcolor, spacing=1.12)


def stat(s, x, y, w, big, label, dark=False, bigc=TERRA, size=34):
    text(s, x, y, w, Inches(0.7), big, size=size, font=DISPLAY, color=bigc)
    text(s, x, y + Inches(0.72), w, Inches(0.7), label, size=11,
         color=FOG if dark else MUTED, spacing=1.1)


def table(s, x, y, colw, rows, rowh=Inches(0.36), size=11, header_fill=INK,
          first_bold=True, hl_col=None, left_all=False):
    """Simple shape-drawn table. rows[0] is the header."""
    yy = y
    for ri, row in enumerate(rows):
        xx = x
        is_h = ri == 0
        if is_h:
            box(s, x, yy, sum(colw), rowh, fill=header_fill)
        elif ri % 2 == 0:
            box(s, x, yy, sum(colw), rowh, fill=CARD)
        for ci, cell in enumerate(row):
            col = BONE if is_h else INK
            b = is_h or (first_bold and ci == 0)
            if hl_col is not None and ci == hl_col and not is_h:
                col, b = TERRA, True
            text(s, xx + Inches(0.1), yy, colw[ci] - Inches(0.15), rowh, str(cell),
                 size=size, color=col, bold=b, anchor=MSO_ANCHOR.MIDDLE,
                 align=PP_ALIGN.LEFT if (ci == 0 or left_all) else PP_ALIGN.RIGHT)
            xx += colw[ci]
        yy += rowh
    box(s, x, yy, sum(colw), Emu(9000), fill=LINE)
    return yy


# ================================================================ SLIDES
# 1 Title -------------------------------------------------------------
s = new_slide(dark=True)
if WORD_BONE.exists():
    image(s, WORD_BONE, MX, Inches(0.8), Inches(3.2), Inches(0.9), align="left")
text(s, MX, Inches(2.2), Inches(6.4), Inches(1.8), "Every street has a story.",
     size=44, font=DISPLAY, color=BONE, spacing=1.0)
text(s, MX, Inches(3.95), Inches(6.0), Inches(1.2),
     "A pocket-sized AI travel companion that tells you the true, sourced stories of the "
     "city around you, through your own earphones. Online or offline. Phone stays in "
     "your pocket.", size=15, color=FOG, spacing=1.2)
text(s, MX, Inches(5.55), Inches(6), Inches(0.4),
     [[("Pre-seed  ·  $1.75M SAFE  ·  $12M post-money cap", {"bold": True, "color": BRASS})]],
     size=14)
text(s, MX, Inches(5.95), Inches(6), Inches(0.4),
     "Nimit Jain, Founder & CEO  ·  LORE Labs, Inc. (Delaware C-corp, planned)  ·  Nov 2026",
     size=11, color=FOG)
box(s, Inches(7.2), Inches(0.6), Inches(5.55), Inches(6.3), fill=BONE)
image(s, A.get("hero-square.png", ""), Inches(7.3), Inches(0.75), Inches(5.35), Inches(3.3))
if "lineup.png" in A:
    image(s, A["lineup.png"], Inches(7.3), Inches(4.0), Inches(5.35), Inches(1.35))
text(s, Inches(7.45), Inches(5.6), Inches(5.1), Inches(1.2),
     [[("Pebble $149  ·  Square $299  ·  Atlas $599", {"bold": True, "color": INK})],
      "Industrial-design renders (design/industrial). Not production units."],
     size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, dark=True)

# 2 Problem -----------------------------------------------------------
s = new_slide()
header(s, "The problem", "In a new city, your phone is your guide, your map, your wallet\nand your only lifeline. It is bad at all four at once.")
cards = [
    ("Eyes on the screen",
     "Travellers read Wikipedia in front of monuments and still miss what is there. "
     "Generic \"top 10\" lists and audio guides tied to single sites."),
    ("Battery dead by 4 pm",
     "Maps, camera, tickets and payments all drain one battery. When it dies you lose "
     "navigation, payments and the ability to call for help."),
    ("Phone out = target",
     "Greater Rome: 33,455 pickpocketing incidents in 2024 (+68% vs 2019). Holding a phone "
     "out marks you as a lost tourist."),
    ("Data is still a hassle",
     ">70% of frequent travellers in one survey had experienced roaming bill shock "
     "(avg. >€200). eSIMs help, but need setup per trip."),
]
cw = Inches(2.92)
for i, (t, b) in enumerate(cards):
    card(s, MX + i * (cw + Inches(0.15)), Inches(2.15), cw, Inches(2.6), t, b, bsize=12)
text(s, MX, Inches(5.1), Inches(12.1), Inches(1.2),
     [[("Who feels it most: ", {"bold": True, "color": INK}),
       ("solo travellers aged 22-45 making 2+ international trips a year. 35-48% of travellers "
        "now plan or take solo trips; solo share of Hostelworld bookings is 63%.", {})]],
     size=14, color=MUTED)
source(s, "Sources: research/trends.md, research/personas.md, research/README.md (quotezone.co.uk pickpocketing aggregation, "
          "telecompaper.com roaming survey, Skyscanner/Global Rescue/IMG). Secondary aggregations; indicative.")

# 3 Solution / the walk ----------------------------------------------
s = new_slide()
header(s, "The solution", "Walk with a brilliant local friend who knows the history.")
steps = [
    ("01", "You walk", "Phone in pocket. LORE clipped on or in hand. Your own earphones (LE Audio / A2DP)."),
    ("02", "LORE knows where you face", "Map Vision fuses L1+L5 GNSS, 9-axis IMU, steps and building outlines."),
    ("03", "It tells you what's there", "\"The building on your right...look up.\" 20-40 s stories from checked, sourced facts."),
    ("04", "You go deeper, or not", "Press: talk. Double-press: what's around me. Long-press: SOS. Quiet while you cross or talk."),
]
sw = Inches(2.92)
for i, (n, t, b) in enumerate(steps):
    x = MX + i * (sw + Inches(0.15))
    text(s, x, Inches(1.85), sw, Inches(0.6), n, size=30, font=DISPLAY, color=TERRA)
    text(s, x, Inches(2.5), sw, Inches(0.5), t, size=15, bold=True)
    text(s, x, Inches(2.95), sw, Inches(1.3), b, size=12, color=MUTED, spacing=1.15)
box(s, MX, Inches(4.35), Inches(12.1), Inches(1.95), fill=INK)
text(s, MX + Inches(0.35), Inches(4.55), Inches(11.4), Inches(1.2),
     [[("“See the Pantheon behind you? Turn around and look up. The building you see was rebuilt "
        "under the emperor Hadrian around 113 to 125 AD. Yet the inscription on the front credits "
        "Marcus Agrippa...”", {"italic": True})]],
     size=16, font=DISPLAY, color=BONE, spacing=1.15)
text(s, MX + Inches(0.35), Inches(5.8), Inches(11.4), Inches(0.4),
     "Actual output of the Rome prototype (simulated walk, offline template path). Facts rome-pantheon-1..4, "
     "source: en.wikipedia.org/wiki/Pantheon,_Rome", size=10, color=BRASS)
source(s, "Source: software/simulations/output/rome-walk-transcript.md (SIMULATED sensors). Principle D4: the AI tells the story; it cannot invent history.")

# 4 Why a device ------------------------------------------------------
s = new_slide()
header(s, "Why a device, not just an app", "Phones are bad at three things travellers need. A device fixes them.")
rows = [
    ("Works when the phone doesn't", "Built-in global eSIM (Square/Atlas, LTE Cat-1bis, 150+ countries). Phone stolen or flat: you still have stories, maps and SOS."),
    ("2+ day battery, separate from the phone", "Simulated travel day: Pebble 2.1, Square 2.1, Atlas 2.4 days (new cell, optimised firmware). Phone battery stays for photos and payments."),
    ("Offline City Packs", "Maps, routing, stories and on-device AI (3B on Square) with no signal. Rome pack: 86 POIs, 215 sourced fact cards today."),
    ("Eyes up, phone hidden", "Physical controls: press to talk, double-press 'what's around me', long-press SOS. Walk-me-home check-ins."),
]
for i, (t, b) in enumerate(rows):
    y = Inches(1.85) + i * Inches(1.05)
    box(s, MX, y, Inches(7.0), Inches(0.92), fill=CARD, line=LINE)
    box(s, MX, y, Inches(0.06), Inches(0.92), fill=TERRA)
    text(s, MX + Inches(0.25), y + Inches(0.1), Inches(6.6), Inches(0.35), t, size=14, bold=True)
    text(s, MX + Inches(0.25), y + Inches(0.42), Inches(6.6), Inches(0.5), b, size=10.5, color=MUTED)
image(s, A.get("exploded-square.png", ""), Inches(7.85), Inches(1.75), Inches(4.9), Inches(3.2))
text(s, Inches(7.85), Inches(5.0), Inches(4.9), Inches(1.5),
     [[("Honest test (Gate G0): ", {"bold": True, "color": INK}),
       ("if the phone app alone delivers ≥80% of the value, we ship the app first and keep the "
        "device as a premium model (decision D1).", {})]],
     size=12, color=MUTED)
source(s, "Sources: COMPANY.md 'why buy'; hardware/simulations/results.md (SIMULATED battery); software/README.md; company/decisions-log.md D1, D11.")

# 5 Product family ----------------------------------------------------
s = new_slide()
header(s, "The product family", "Good / better / best, one software platform.")
fam = [
    ("hero-pebble.png", "LORE Pebble", "$149", "Clip-on, 42×42×14 mm, LED ring. 'Lite' design (D10): AI via phone over BLE; offline pre-generated stories. 1.5-day target."),
    ("hero-square.png", "LORE Square", "$299", "54×54×12 mm aluminium, 1.6\" AMOLED, crown. LTE Cat-1bis + global eSIM (D11), on-device 3B model. 2-day target."),
    ("hero-atlas.png", "LORE Atlas", "$599", "Titanium + sapphire, 13.5 mm (D12), Glance camera, satellite SOS (NTN). On-device 3-4B + vision. '2+ days'."),
]
fw = Inches(3.95)
for i, (img, name, price, desc) in enumerate(fam):
    x = MX + i * (fw + Inches(0.12))
    box(s, x, Inches(1.75), fw, Inches(2.75), fill=RGBColor(0xEE, 0xE5, 0xD8))
    image(s, A.get(img, ""), x, Inches(1.8), fw, Inches(2.65))
    text(s, x, Inches(4.6), fw * 0.65, Inches(0.4), name, size=17, font=DISPLAY)
    text(s, x + fw * 0.6, Inches(4.6), fw * 0.4, Inches(0.4), price, size=17, font=DISPLAY,
         color=TERRA, align=PP_ALIGN.RIGHT)
    text(s, x, Inches(5.05), fw, Inches(1.3), desc, size=11, color=MUTED, spacing=1.15)
text(s, MX, Inches(6.0), Inches(12.1), Inches(0.4),
     [[("LORE Pass: ", {"bold": True, "color": INK}),
       ("Explorer $9.99/mo or $79/yr · Trip Pass $29 / 14 days / 5 GB · $49 / 30 days / 10 GB · free tier (3 cities). "
        "Every device includes 3 months of Explorer.  ", {}),
       ("LORE Kit ", {"bold": True, "color": INK}), ("accessories incl. LORE Buds $129.", {})]],
     size=11.5, color=MUTED)
source(s, "Renders: design/industrial/renders (concept renders, not production units). Specs: COMPANY.md, hardware/README.md, decisions D8, D10-D12.")

# 6 How Map Vision works ----------------------------------------------
s = new_slide()
header(s, "How Map Vision works", "LORE 'sees' without a camera: sensor fusion + map geometry.")
inputs = ["Multi-band GNSS (L1+L5)", "9-axis IMU + compass", "Step & walking direction", "Building footprints, street sides, entrances", "Places + live data"]
for i, t in enumerate(inputs):
    y = Inches(1.8) + i * Inches(0.62)
    box(s, MX, y, Inches(3.3), Inches(0.5), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, MX + Inches(0.15), y, Inches(3.1), Inches(0.5), t, size=11.5, anchor=MSO_ANCHOR.MIDDLE)
    arr = box(s, MX + Inches(3.4), y + Inches(0.2), Inches(0.5), Inches(0.1), fill=BRASS)
mid_x = MX + Inches(4.05)
box(s, mid_x, Inches(1.8), Inches(2.9), Inches(2.98), fill=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, mid_x + Inches(0.2), Inches(1.95), Inches(2.5), Inches(2.7),
     [[("Heading fusion", {"bold": True, "color": BONE, "size": 15})],
      "Kalman filter: heading + compass bias; rejects magnetic disturbance.",
      [("Story engine", {"bold": True, "color": BONE, "size": 15})],
      "Interest, distance, visibility, novelty, cooldowns, quiet moments (crossing / talking).",
      [("Fact guard", {"bold": True, "color": BONE, "size": 15})],
      "Narration may use ONLY the selected fact cards; else a deterministic template."],
     size=10.5, color=FOG, spacing=1.1, space_after=3)
box(s, mid_x + Inches(3.0), Inches(3.2), Inches(0.5), Inches(0.1), fill=BRASS)
text(s, mid_x + Inches(3.6), Inches(1.85), Inches(4.5), Inches(1.2),
     "left · right · ahead · behind · look up", size=18, font=DISPLAY, color=TERRA)
text(s, mid_x + Inches(3.6), Inches(2.35), Inches(4.5), Inches(2.4),
     [[("Simulated accuracy (hardware/simulations): ", {"bold": True, "color": INK}),
       ("left/right 99.9% with EKF + map-matching + L1/L5; EKF alone 98.9%; L1+L5 adds 3-7 points over L1-only.", {})],
      [("Prototype walk (software sim): ", {"bold": True, "color": INK}),
       ("direction 100% vs ground truth; fused heading error mean 4.3°, p95 12.7°.", {})],
      [("Gate G1 bar on real walks: ", {"bold": True, "color": INK}), ("≥90%. D3 kill line: <85%.", {})]],
     size=11.5, color=MUTED, spacing=1.15, space_after=6)
box(s, MX, Inches(5.1), Inches(12.1), Inches(1.3), fill=CARD, line=LINE)
text(s, MX + Inches(0.25), Inches(5.22), Inches(11.6), Inches(1.1),
     [[("Why no camera on the first models (D3): ", {"bold": True, "color": INK}),
       ("avoids privacy problems, cost and battery drain. Only Atlas adds an opt-in 'Glance' camera. "
        "Patent provisionals planned before any public demo: P1 Map Vision, P2 quiet-moment (legal-ops/ip-strategy.md).", {})],
      [("Caveat: ", {"bold": True, "color": TERRA}),
       ("all accuracy figures are SIMULATED. Real-walk validation in Rome starts with the EVT beta (Q1 2027).", {})]],
     size=11.5, color=MUTED)
source(s, "Sources: hardware/README.md, hardware/simulations/results.md, software/packages/core (heading-fusion.ts, story-engine.ts, narration.ts).")

# 7 Demo ---------------------------------------------------------------
s = new_slide()
header(s, "Demo: the working Rome prototype", "The real story engine, walking Pantheon to Trevi (1.93 km).")
shots = ["shot-pantheon.png", "shot-navona.png", "shot-around.png", "shot-dark.png"]
sx = MX
for sh in shots:
    image(s, A.get(sh, ""), sx, Inches(1.75), Inches(1.6), Inches(3.3))
    sx += Inches(1.72)
text(s, MX, Inches(5.15), Inches(6.8), Inches(0.6),
     "Simulation mode, phone viewport: Pantheon · Piazza Navona · 'What's around me' at Largo Argentina · dark 'storyteller' persona.",
     size=10, color=MUTED)
mx2 = Inches(7.7)
text(s, mx2, Inches(1.75), Inches(5), Inches(0.4), "Walk-simulation metrics (SIMULATED)", size=13, bold=True)
rows = [["Metric", "Value"],
        ["Narrations", "17 (8.8 / km)"],
        ["Facts spoken / repeats", "39 / 0.0%"],
        ["Quiet-moment compliance", "100% (0 violations)"],
        ["Direction accuracy vs truth", "100%"],
        ["Fused heading error", "4.3° mean · 12.7° p95"],
        ["Share of time speaking", "26%"],
        ["Rome City Pack", "86 POIs · 215 fact cards"],
        ["Tests passing", "61 vitest + 16 pytest"]]
table(s, mx2, Inches(2.2), [Inches(2.85), Inches(2.2)], rows, rowh=Inches(0.37), size=11)
text(s, MX, Inches(5.75), Inches(12.1), Inches(0.8),
     [[("What is not proven yet: ", {"bold": True, "color": TERRA}),
       ("no real-user walks, no EVT hardware. Seed facts are curated with Wikipedia citations but still need an online verify pass "
        "and editorial review before customers hear them. Gate G0 (Nov 2026-Jan 2027): 20 Rome walk tests + 200 remote sim users.", {})]],
     size=11.5, color=MUTED)
source(s, "Sources: software/README.md, software/apps/prototype/screenshots/, software/simulations/output/rome-walk-transcript.md.")

# 8 Market ---------------------------------------------------------------
s = new_slide()
header(s, "Market", "A $10.7B/yr job market today; we size the beachhead bottom-up.")
stats = [("$10.7B", "TAM (top-down, 2027): eSIM, audio guides, travel safety, AI-companion, self-guided tours"),
         ("$24.0B", "TAM (bottom-up): 145M frequent international travellers × $165 ARPU"),
         ("$1.44B", "SAM: 8.7M solo/independent travellers, 2+ trips/yr, US/CA/UK/EU/AU-NZ, will pay $150+"),
         ("$93M", "SOM 2030 base: 250k units in 2030, 415k cumulative (4.7% of SAM). Bear $46M / Bull $148M")]
for i, (b, l) in enumerate(stats):
    stat(s, MX, Inches(1.7) + i * Inches(1.15), Inches(4.2), b, l, size=30)
box(s, Inches(5.1), Inches(1.75), Inches(7.6), Inches(3.75), fill=WHITE, line=LINE)
image(s, RCH / "market-size-waterfall.png", Inches(5.2), Inches(1.85), Inches(7.4), Inches(3.55))
text(s, Inches(5.1), Inches(5.65), Inches(7.6), Inches(0.9),
     "Tailwinds: 1.52B international arrivals in 2025 (+4%); top-100 city arrivals +8%; travel eSIM $1.8B (2025, +85%) → $8.7B (2030); "
     "solo-travel spend ~$482B (2024) → ~$1.07T (2030).", size=11, color=MUTED)
source(s, "Source: research/market-sizing.md (all figures cited with URLs there), research/charts/market-size-waterfall.png. SOM is a Research estimate [EST].")

# 9 Competition -------------------------------------------------------
s = new_slide()
header(s, "Competition and positioning", "Nothing else is standalone, travel-first and verified.")
box(s, MX, Inches(1.7), Inches(5.6), Inches(4.5), fill=WHITE, line=LINE)
image(s, RCH / "positioning-map.png", MX + Inches(0.1), Inches(1.75), Inches(5.4), Inches(4.4))
comp = [["Threat", "What it is", "Our answer"],
        ["Google/Samsung Gemini glasses + Maps", "Fall 2026; camera + assistant", "Travel brain; works with any audio"],
        ["Meta Ray-Ban / Display", "~7M glasses sold 2025", "No camera needed; $149 entry"],
        ["Phone + AI-guide app + eSIM", "Audiala $4.99/mo + Airalo", "Battery, theft, eyes-up, SOS"],
        ["OpenAI io device", "H2 2026-27 (narrative)", "Narrow job; verified facts"]]
table(s, Inches(6.45), Inches(1.7), [Inches(2.35), Inches(1.85), Inches(2.1)], comp, rowh=Inches(0.55), size=9.5, left_all=True)
text(s, Inches(6.45), Inches(4.65), Inches(6.3), Inches(1.6),
     [[("Lessons we design against: ", {"bold": True, "color": INK}),
       ("Humane (~10k units; assets sold to HP for $116M) and Rabbit (~130k sold, ~5% daily use) shipped general assistants with no "
        "daily job. Plaud (>1M devices, ~$250M run-rate) shows one narrow job + clip-on + subscription works.", {})],
      [("Moat we are building: ", {"bold": True, "color": INK}),
       ("sourced City Packs and the fact-checking process; Map Vision IP; brand.", {})]],
     size=11, color=MUTED, spacing=1.12)
source(s, "Source: research/competitors.md (30-product matrix, URLs inside), research/charts/positioning-map.png.")

# 10 Business model & pricing -------------------------------------------
s = new_slide()
header(s, "Business model and pricing", "Hardware gets us in; LORE Pass and bookings compound.")
pr = [["", "Pebble", "Square", "Atlas"],
      ["US price", "$149", "$299", "$599"],
      ["Kickstarter super early bird", "$99", "$199", "$399"],
      ["EU / UK", "€159 / £149", "€329 / £299", "€649 / £599"],
      ["Price anchor", "Plaud NotePin S $179", "Pocketalk Plus $299.95", "inReach Mini 3+ $499 / Watch Ultra $799"]]
table(s, MX, Inches(1.75), [Inches(2.6), Inches(1.75), Inches(1.75), Inches(2.4)], pr, rowh=Inches(0.48), size=10.5)
streams = [("Hardware", "DTC first; retail after returns <5%, led by Pebble + Atlas (D15)"),
           ("LORE Pass", "$9.99/mo · $79/yr · Trip Pass $29 / $49; 3 months free with every device"),
           ("LORE Kit", "Clip, lanyard, dock, case, LORE Buds $129, city bundles"),
           ("Bookings", "Affiliate fees: hotels, tours, restaurants"),
           ("B2B fleet", "Rental for hostels and tour operators (pilots Q4 2027+)")]
for i, (t, b) in enumerate(streams):
    y = Inches(1.75) + i * Inches(0.86)
    box(s, Inches(9.35), y, Inches(3.38), Inches(0.76), fill=CARD, line=LINE)
    box(s, Inches(9.35), y, Inches(0.06), Inches(0.76), fill=[TERRA, BRASS, OLIVE, FOG, INK][i])
    text(s, Inches(9.55), y + Inches(0.06), Inches(3.1), Inches(0.3), t, size=12, bold=True)
    text(s, Inches(9.55), y + Inches(0.33), Inches(3.1), Inches(0.45), b, size=9.5, color=MUTED)
text(s, MX, Inches(4.35), Inches(8.5), Inches(1.9),
     [[("Why these prices: ", {"bold": True, "color": INK}),
       ("Research's comparable-price ladder (research/pricing-research.md). Atlas raised $549 → $599 (D8), to be A/B tested; "
        "revert if the survey shows >20% loss of purchase intent.", {})],
      [("Retention by design (the Rabbit lesson): ", {"bold": True, "color": INK}),
       ("annual Explorer at $79, a free tier that stays useful, home-city 'Weekend Walks', and retention KPIs reported to investors.", {})]],
     size=11.5, color=MUTED, spacing=1.15, space_after=8)
source(s, "Sources: COMPANY.md, research/pricing-research.md, company/decisions-log.md D6, D8, D15.")

# 11 Unit economics -----------------------------------------------------
s = new_slide()
header(s, "Unit economics", "We found the margin problem early and changed the product.")
ue = [["Landed COGS @10k units", "Before (full spec)", "Decision", "After"],
      ["Pebble $149", "$157 incl. NRE → −5% GM", "D10: 'Lite' (AI via phone)", "~$90 → 39% GM"],
      ["Square $299", "$299 incl. NRE → 0% GM", "D11: LTE Cat-1bis, cost-down", "Target ≤$240 by unit 10k; $329 fallback"],
      ["Atlas $599", "$413 incl. NRE → 25% at $549", "D8 price $599; D12 13.5 mm", "Margin engine"],
      ["Any SKU, China FATP", "Square $343 landed (tariffs)", "D13: Vietnam / Malaysia", "0% duty assumed"]]
table(s, MX, Inches(1.75), [Inches(2.35), Inches(2.85), Inches(2.85), Inches(4.05)], ue, rowh=Inches(0.48), size=10.5, hl_col=3, left_all=True)
uex = FIN["unit_economics"] if FIN else None
if FCH.joinpath("unit_economics_waterfall.png").exists():
    box(s, MX, Inches(4.3), Inches(7.3), Inches(2.3), fill=WHITE, line=LINE)
    image(s, FCH / "unit_economics_waterfall.png", MX + Inches(0.05), Inches(4.35), Inches(7.2), Inches(2.2))
text(s, Inches(8.1), Inches(4.3), Inches(4.65), Inches(2.3), "UE_PLACEHOLDER", size=11, color=MUTED)
UE_BOX = s.shapes[-1]
source(s, "Sources: hardware/README.md + hardware/bom/BOM.md (estimates, 2026 DRAM prices), finance/unit-economics.md, company/decisions-log.md D10-D13.")

# 12 GTM & Kickstarter ---------------------------------------------------
s = new_slide()
header(s, "Go-to-market and Kickstarter", "Stories sell; safety justifies; Kickstarter proves demand.")
ph = [("Nov 26-Jan 27", "Waitlist", "5k sign-ups"), ("Feb-Apr 27", "Rome beta", "50 EVT units"),
      ("Pre-launch", "List + VIPs", "20k emails, 2k $1 VIPs"), ("25 May-24 Jun 27", "Kickstarter", "$250k goal"),
      ("Q4 27", "DTC", "First 5k units"), ("2028", "Retail / airport", "after returns <5%")]
pw = Inches(1.95)
for i, (when, what, kpi) in enumerate(ph):
    x = MX + i * (pw + Inches(0.08))
    box(s, x, Inches(1.8), pw, Inches(1.25), fill=INK if i == 3 else CARD, line=None if i == 3 else LINE)
    text(s, x + Inches(0.15), Inches(1.9), pw - Inches(0.25), Inches(0.3), when, size=9.5, color=BRASS, bold=True)
    text(s, x + Inches(0.15), Inches(2.18), pw - Inches(0.25), Inches(0.4), what, size=14, bold=True, color=BONE if i == 3 else INK)
    text(s, x + Inches(0.15), Inches(2.58), pw - Inches(0.25), Inches(0.4), kpi, size=10, color=FOG if i == 3 else MUTED)
stat(s, MX, Inches(3.35), Inches(2.8), "$1.04M", "Kickstarter base case: 4,600 backers × $225 avg pledge (plan, not a result)")
stat(s, MX + Inches(3.05), Inches(3.35), Inches(2.8), "≈$50", "Cost per backer on $230k marketing (Gate G2 needs <$60)")
stat(s, MX + Inches(6.1), Inches(3.35), Inches(2.8), "$554k", "Conservative case: misses $1M; escalation playbook ready")
stat(s, MX + Inches(9.15), Inches(3.35), Inches(2.9), "$70", "DTC blended CAC ceiling after launch")
text(s, MX, Inches(4.9), Inches(12.1), Inches(1.6),
     [[("Channels: ", {"bold": True, "color": INK}),
       ("TikTok/IG organic ('Secrets of Rome') → paid lead gen ($2.75 CPL target) → creators (≤$45 CPB) → PR (15+ tier-1/2 articles) → email (8% list conversion).", {})],
      [("Messaging: ", {"bold": True, "color": INK}), ("Look up · True stories · Go anywhere, safely. Lead with experience, justify with safety.", {})],
      [("Honesty: ", {"bold": True, "color": TERRA}), ("no partners, customers or pre-orders exist yet; all named companies are targets.", {})]],
     size=11.5, color=MUTED, spacing=1.12, space_after=6)
source(s, "Sources: gtm/README.md, gtm/gtm-strategy.md, gtm/kickstarter/, finance/kickstarter-economics.md.")

# 13 Roadmap & gates -------------------------------------------------
s = new_slide()
header(s, "Roadmap, milestones and gates", "Each stage removes the biggest risk before we spend on the next.")
gates = [("Stage 0 · Nov 26-Jan 27", "Prove the magic", "Phone PWA + Rome pack (300+ places hand-checked); 20 Rome walks; n≥500 survey",
          "G0: ≥40% 'very disappointed'; ≥60% finish a 30-min walk; ≥25% would pay $149+"),
         ("Stage 1 · Feb-May 27", "Prove the device", "50 EVT units; ODM selected (VN/MY); SoC locked; pre-seed closes; first hires",
          "G1: battery ≥1.5 days measured; left/right ≥90% on walks; beta use ≥3 days/week"),
         ("Stage 2 · Jun-Aug 27", "Prove demand", "Kickstarter (25 May-24 Jun); 20k list; DVT; certification starts",
          "G2: ≥$1M pledged; ≥3,000 backers; CPB <$60"),
         ("Stage 3 · Sep 27-Apr 28", "Ship", "$6M seed; PVT → first 5k units; FCC/CE/UKCA/carriers; 1 → 8 cities; B2B pilots",
          "G3: returns <5%; attach >55%; churn <6%/mo; NPS >50")]
gw = Inches(2.95)
for i, (when, title, work, gate) in enumerate(gates):
    x = MX + i * (gw + Inches(0.1))
    box(s, x, Inches(1.8), gw, Inches(0.08), fill=[TERRA, BRASS, OLIVE, INK][i])
    text(s, x, Inches(1.98), gw, Inches(0.3), when, size=10, bold=True, color=TERRA)
    text(s, x, Inches(2.28), gw, Inches(0.45), title, size=19, font=DISPLAY)
    text(s, x, Inches(2.82), gw, Inches(1.4), work, size=11, color=MUTED, spacing=1.15)
    box(s, x, Inches(3.85), gw, Inches(1.45), fill=CARD, line=LINE)
    text(s, x + Inches(0.15), Inches(3.95), gw - Inches(0.3), Inches(1.3), gate, size=11, bold=True, spacing=1.15)
text(s, MX, Inches(5.6), Inches(12.1), Inches(0.6),
     [[("What we will not do: ", {"bold": True, "color": INK}),
       ("no glasses; no camera before Atlas; ≤1 new city/month until City Pack quality is measured; no retail before DTC returns <5%. "
        "2028: retail + 20 cities.", {})]], size=11.5, color=MUTED)
source(s, "Source: company/operating-plan.md, COMPANY.md timeline, gtm/README.md.")

# 14 Team & hiring ---------------------------------------------------------
s = new_slide()
header(s, "Team and hiring", "Founder-led today; the pre-seed buys the founding team.")
card(s, MX, Inches(1.8), Inches(4.0), Inches(2.3), "Nimit Jain, Founder & CEO",
     "Product, fundraising and story quality. Today every department (research, brand, design, hardware, software, "
     "finance, GTM, legal/ops, IR) is run as an AI-agent team with founder sign-off on every output; this deck is built from their files.",
     bsize=11)
card(s, MX, Inches(4.25), Inches(4.0), Inches(2.1), "Contract-first, to de-risk",
     "Industrial-design studio, EE/RF consultancy, ODM programme manager, fractional CFO and GC, patent/TM attorneys. "
     "Pre-seed team 6-8 people in 2027; ~20 at seed (2028).", bsize=11, accent=BRASS)
hires = [["#", "First hires (pre-seed)", "When"],
         ["1", "Founding Embedded / Firmware Lead", "Nov 2026"],
         ["2", "Founding ML / Story-Engine Engineer", "Nov 2026"],
         ["3", "Mobile + Backend Engineer", "Dec 2026"],
         ["4", "Hardware / Electrical Lead (contract first)", "Jan 2027"],
         ["5", "Sensor-Fusion / Map Vision Engineer", "Jan 2027"],
         ["6", "Content Editor-in-Chief (City Packs)", "Jan 2027"],
         ["7", "Growth & Community Lead (Kickstarter)", "Feb 2027"],
         ["8", "Supply Chain & Quality Manager", "Mar 2027"]]
table(s, Inches(4.85), Inches(1.8), [Inches(0.45), Inches(5.6), Inches(1.85)], hires, rowh=Inches(0.42), size=11, left_all=True)
text(s, Inches(4.85), Inches(5.7), Inches(7.9), Inches(0.7),
     "We are not claiming hires, advisors or investors that do not exist. Roles, salary bands (US/EU/IN) and equity bands: legal-ops/hr-hiring.md. 12% option pool planned.",
     size=10.5, color=MUTED)
source(s, "Sources: company/org-chart.md, legal-ops/hr-hiring.md, legal-ops/README.md.")

# 15 Financials ---------------------------------------------------------------
s = new_slide()
header(s, "Financials", "FIN_TITLE")
FIN_SLIDE = s
# 16 Ask + contact -----------------------------------------------------------
s = new_slide(dark=True)
ASK_SLIDE = s

# ================================================================ finance-dependent
FIN_STATUS = os.environ.get("LORE_FIN_STATUS", "final" if (ROOT / "finance/README.md").exists() else "draft")


def fill_finance():
    s = FIN_SLIDE
    title_shape = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == "FIN_TITLE"][0]
    if FIN:
        b = FIN["scenario_snapshot"]["Base"]
        fy = FIN["fy"]
        title_shape.text_frame.paragraphs[0].runs[0].text = (
            f"Base case: {money(b['rev'][-1], 0)} revenue and EBITDA-positive in FY{b['be']}.")
        rows = [["Base case", *[f"FY{str(y)[2:]}" for y in fy[1:]]],
                ["Units", *[f"{u/1000:.1f}k" if u < 1e4 else f"{u/1000:.0f}k" for u in b["units"][1:]]],
                ["Revenue", *[money(v) for v in b["rev"][1:]]],
                ["Gross margin", *[f"{g*100:.0f}%" for g in b["gm"][1:]]],
                ["EBITDA", *[money(v) for v in b["ebitda"][1:]]],
                ["Subscription ARR", *[money(v) for v in b["arr"][1:]]]]
        table(s, MX, Inches(1.75), [Inches(1.75)] + [Inches(0.98)] * 5, rows, rowh=Inches(0.42), size=10.5)
        bear, bull = FIN["scenario_snapshot"]["Bear"], FIN["scenario_snapshot"]["Bull"]
        f = FIN["funding"]
        notes = [[("Scenarios FY31 revenue: ", {"bold": True, "color": INK}),
                  (f"Bear {money(bear['rev'][-1])} · Base {money(b['rev'][-1])} · Bull {money(bull['rev'][-1])}.", {})],
                 [("Funding need: ", {"bold": True, "color": INK}),
                  (f"{money(f['Total equity raised through FY2031'])} total equity to break-even (D16): pre-seed → $6M seed → $3M seed-2 "
                   "→ $25M A → $20M B. Bear case needs ≈$72M. "
                   f"At $1.5M the pre-seed bottomed at {money(f['Minimum month-end cash, Nov-26..Dec-28'])} in Jun-27, "
                   "so the round is now $1.75M (D14).", {})],
                 [("Research cross-check: ", {"bold": True, "color": INK}),
                  ("SOM 2030 base $93M (bear $46M / bull $148M).", {})]]
        text(s, MX, Inches(4.45), Inches(6.7), Inches(2.0), notes, size=11, color=MUTED, spacing=1.12, space_after=6)
        chart = FCH / "revenue_by_stream.png"
        box(s, Inches(7.55), Inches(1.75), Inches(5.2), Inches(2.4), fill=WHITE, line=LINE)
        image(s, chart, Inches(7.6), Inches(1.8), Inches(5.1), Inches(2.3))
        box(s, Inches(7.55), Inches(4.25), Inches(5.2), Inches(2.3), fill=WHITE, line=LINE)
        image(s, FCH / "cash_runway.png", Inches(7.6), Inches(4.3), Inches(5.1), Inches(2.2))
        tag = "" if FIN_STATUS == "final" else "  [DRAFT: finance/README.md not final at build time]"
        source(s, "Source: finance/LORE_Financial_Model.xlsx via finance/model_summary.json; finance/charts. All projections are estimates." + tag)
    else:
        title_shape.text_frame.paragraphs[0].runs[0].text = "[PLACEHOLDER] Financial model pending; Research SOM shown."
        stat(s, MX, Inches(2.0), Inches(3.5), "$93M", "SOM 2030 base (Research) [PLACEHOLDER for model]")
        stat(s, MX + Inches(4), Inches(2.0), Inches(3.5), "250k", "Units in 2030 (Research SOM)")
        stat(s, MX + Inches(8), Inches(2.0), Inches(3.5), "$46M-$148M", "Bear-Bull 2030 revenue (Research)")
        source(s, "[PLACEHOLDER] finance/ not available at build time. Source: research/market-sizing.md.")


def fill_ask(use_of_funds, runway_note):
    s = ASK_SLIDE
    text(s, MX, Inches(0.42), Inches(9), Inches(0.3), "THE ASK", size=11, color=TERRA, bold=True)
    text(s, MX, Inches(0.72), Inches(7), Inches(1.4), "$1.75M pre-seed SAFE\n$12M post-money cap",
         size=34, font=DISPLAY, color=BONE, spacing=1.0)
    text(s, MX, Inches(2.15), Inches(6.2), Inches(1.6),
         [[("Gets us to: ", {"bold": True, "color": BONE}),
           ("Gate G0 (product love), 50 EVT units and Gate G1, a $1M+ Kickstarter (Gate G2), and the $6M seed.", {})],
          [("Investor owns: ", {"bold": True, "color": BONE}), ("≈14.6% at the $12M post-money cap. $1.5M + $250k top-up (D14).", {})],
          [(runway_note, {})]],
         size=12.5, color=FOG, spacing=1.15, space_after=6)
    text(s, Inches(7.3), Inches(0.75), Inches(5.4), Inches(0.4), "Use of funds", size=16, bold=True, color=BONE)
    y = Inches(1.2)
    total = sum(v for _, v, _ in use_of_funds)
    for i, (label, v, note) in enumerate(use_of_funds):
        frac = v / total if total else 0
        text(s, Inches(7.3), y, Inches(3.8), Inches(0.3), label, size=12, color=BONE, bold=True)
        text(s, Inches(11.1), y, Inches(1.6), Inches(0.3), f"{money(v)} · {frac*100:.0f}%", size=12, color=BRASS,
             align=PP_ALIGN.RIGHT)
        box(s, Inches(7.3), y + Inches(0.33), Inches(5.4), Inches(0.09), fill=RGBColor(0x2A, 0x2D, 0x33))
        box(s, Inches(7.3), y + Inches(0.33), max(int(Inches(5.4) * frac), Inches(0.05)), Inches(0.09),
            fill=[TERRA, BRASS, OLIVE, FOG, BONE, TERRA][i % 6])
        text(s, Inches(7.3), y + Inches(0.45), Inches(5.4), Inches(0.25), note, size=9, color=FOG)
        y += Inches(0.68)
    # contact strip
    box(s, Inches(0), Inches(5.35), W, Inches(1.55), fill=TERRA)
    text(s, MX, Inches(5.55), Inches(6.5), Inches(0.6), "Let's walk Rome together.", size=26, font=DISPLAY, color=BONE)
    text(s, MX, Inches(6.2), Inches(7), Inches(0.5),
         "Live prototype demo on request. Data room index: investor/data-room-index.md", size=12, color=BONE)
    text(s, Inches(7.3), Inches(5.55), Inches(5.4), Inches(1.2),
         [[("Nimit Jain", {"bold": True, "size": 15})], "Founder & CEO, LORE Labs, Inc.",
          "[founder email]  ·  lore.travel (domain to be confirmed)"],
         size=12, color=BONE, spacing=1.15, space_after=2)
    footer(s, dark=True)


def fill_ue_box(lines):
    tf = UE_BOX.text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    first = True
    for head, body in lines:
        if not first:
            p = tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        p.line_spacing = 1.12
        for t, b, c in [(head, True, INK), (body, False, MUTED)]:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(11)
            r.font.bold = b
            r.font.name = BODY
            r.font.color.rgb = c


# ---- values pulled from finance docs (update when finance/README.md is final) ----
from fin_inputs import USE_OF_FUNDS, RUNWAY_NOTE, UE_LINES  # noqa: E402

fill_finance()
fill_ask(USE_OF_FUNDS, RUNWAY_NOTE)
fill_ue_box(UE_LINES)

prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides); finance status: {FIN_STATUS}")
